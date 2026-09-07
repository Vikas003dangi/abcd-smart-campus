"""
users/utils/floor_export.py
Generates in-memory PDF and Excel (.xlsx) data sheets for Library Floor Status.
Directly streams to browser via io.BytesIO without storing temporary files on disk.
"""

import io
from datetime import datetime
from django.utils import timezone
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfgen import canvas
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from users.models import Seat, FeeTransaction


def get_floor_export_data(floor_name):
    """
    Fetches and prepares all seat and student fee status data for the specified floor.
    """
    seats = list(
        Seat.objects.filter(floor=floor_name)
        .prefetch_related('assignments__student', 'special_requests__student')
    )
    # Sort seats numerically
    seats.sort(key=lambda s: int(s.seat_number) if str(s.seat_number).isdigit() else 9999)

    today = timezone.localdate()
    export_rows = []

    counts = {
        'total': len(seats),
        'available': 0,
        'occupied': 0,
        'on_hold': 0,
        'shift_occupied': 0,
        'temporary': 0,
        'locked': 0,
    }

    for seat in seats:
        active_assignments = [a for a in seat.assignments.all() if a.is_active]
        is_locked = seat.is_locked

        if is_locked:
            counts['locked'] += 1

        if active_assignments:
            is_shift_occ = len(active_assignments) == 1 and seat.is_shift_enabled

            for a in active_assignments:
                st = a.student
                # Get latest fee transaction
                last_txn = FeeTransaction.objects.filter(student=st).order_by('-payment_date', '-created_at').first()
                if last_txn and last_txn.total_amount:
                    last_amount = f"Rs. {last_txn.total_amount:,.0f}"
                else:
                    last_amount = "—"

                # Fee expiry
                if last_txn and last_txn.expiry_date:
                    fee_exp = last_txn.expiry_date.strftime('%d/%m/%Y')
                elif st.fee_expiry_date:
                    fee_exp = st.fee_expiry_date.strftime('%d/%m/%Y')
                else:
                    fee_exp = "—"

                # Hold expiry
                if a.hold_status == 'active' and a.hold_end_date:
                    hold_exp = a.hold_end_date.strftime('%d/%m/%Y')
                    status_label = "On Hold"
                    counts['on_hold'] += 1
                elif a.is_partial:
                    hold_exp = "—"
                    status_label = "Temporary"
                    counts['temporary'] += 1
                elif is_shift_occ:
                    hold_exp = "—"
                    status_label = "Shift Occupied"
                    counts['shift_occupied'] += 1
                else:
                    hold_exp = "—"
                    status_label = "Occupied"
                    counts['occupied'] += 1

                shift_label = a.shift_type.capitalize() if a.shift_type else ("Shift" if seat.is_shift_enabled else "Full Day")

                export_rows.append({
                    'seat_number': seat.seat_number,
                    'shift': shift_label,
                    'student_name': st.full_name or st.user.username,
                    'mobile': st.mobile_number or "—",
                    'status': status_label,
                    'last_amount': last_amount,
                    'fee_expiry': fee_exp,
                    'hold_expiry': hold_exp,
                })
        else:
            # Vacant seat or locked/hold at seat-level
            if seat.status == 'on_hold':
                status_label = "On Hold"
                counts['on_hold'] += 1
                hold_exp = seat.hold_end_date.strftime('%d/%m/%Y') if seat.hold_end_date else "—"
            elif is_locked:
                status_label = "Locked"
                hold_exp = "—"
            else:
                status_label = "Available"
                counts['available'] += 1
                hold_exp = "—"

            shift_label = "Morning & Evening" if seat.is_shift_enabled else "Full Day"

            export_rows.append({
                'seat_number': seat.seat_number,
                'shift': shift_label,
                'student_name': "—",
                'mobile': "—",
                'status': status_label,
                'last_amount': "—",
                'fee_expiry': "—",
                'hold_expiry': hold_exp,
            })

    return {
        'floor': floor_name,
        'rows': export_rows,
        'counts': counts,
        'date_str': today.strftime('%d/%m/%Y')
    }


class NumberedCanvas(canvas.Canvas):
    """Adds professional page numbers and running footer to ReportLab PDFs."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_footer(num_pages)
            super().showPage()
        super().save()

    def draw_page_footer(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        footer_text = f"Page {self._pageNumber} of {page_count}  •  ABCD Smart Campus Official Library Register  •  Confidential"
        self.drawCentredString(A4[1] / 2, 20, footer_text)
        self.restoreState()


def generate_floor_data_pdf(floor_name):
    """
    Builds a styled A4 landscape PDF document in memory.
    """
    data = get_floor_export_data(floor_name)
    rows = data['rows']
    counts = data['counts']
    date_str = data['date_str']

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#0f172a'),
    )
    meta_style = ParagraphStyle(
        'MetaStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#475569'),
    )
    badge_style = ParagraphStyle(
        'BadgeStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1e293b'),
    )
    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1e293b'),
    )
    cell_bold = ParagraphStyle(
        'TableCellBold',
        parent=cell_style,
        fontName='Helvetica-Bold',
    )
    th_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        alignment=1, # Center
    )

    elements = []

    # Title & Header
    title_text = f"ABCD Smart Library • {floor_name} Data Sheet"
    sub_text = f"Generated on: <b>{date_str}</b>  |  Total Seats: <b>{counts['total']}</b>"
    elements.append(Paragraph(title_text, title_style))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph(sub_text, meta_style))
    elements.append(Spacer(1, 8))

    # Summary KPI Stats Row
    summary_text = (
        f"🟢 Available: <b>{counts['available']}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"🔵 Occupied: <b>{counts['occupied']}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"🟣 Shift Occupied: <b>{counts['shift_occupied']}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"🟠 On Hold: <b>{counts['on_hold']}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"🌸 Temporary: <b>{counts['temporary']}</b> &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"🔒 Locked: <b>{counts['locked']}</b>"
    )
    elements.append(Paragraph(summary_text, badge_style))
    elements.append(Spacer(1, 12))

    # Table Header & Rows
    table_data = [[
        Paragraph("Seat #", th_style),
        Paragraph("Shift / Slot", th_style),
        Paragraph("Student Name", th_style),
        Paragraph("Contact No.", th_style),
        Paragraph("Status", th_style),
        Paragraph("Last Paid", th_style),
        Paragraph("Fee Expiry", th_style),
        Paragraph("Hold Expiry", th_style),
    ]]

    for r in rows:
        # Determine status text color
        stat = r['status']
        if stat == 'Available':
            stat_color = '#059669'
        elif stat in ('Occupied', 'Shift Occupied'):
            stat_color = '#2563eb'
        elif stat == 'On Hold':
            stat_color = '#d97706'
        elif stat == 'Temporary':
            stat_color = '#db2777'
        else:
            stat_color = '#64748b'

        status_cell = Paragraph(f"<font color='{stat_color}'><b>{stat}</b></font>", cell_style)

        table_data.append([
            Paragraph(f"<b>Seat {r['seat_number']}</b>", cell_bold),
            Paragraph(r['shift'], cell_style),
            Paragraph(r['student_name'], cell_style),
            Paragraph(r['mobile'], cell_style),
            status_cell,
            Paragraph(r['last_amount'], cell_style),
            Paragraph(r['fee_expiry'], cell_style),
            Paragraph(r['hold_expiry'], cell_style),
        ])

    # Table styling (A4 landscape usable width approx 794 pt)
    col_widths = [55, 95, 175, 85, 95, 80, 75, 75]
    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e293b')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
    ]))

    elements.append(table)

    doc.build(elements, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


def generate_floor_data_excel(floor_name):
    """
    Builds a styled, formatted Microsoft Excel (.xlsx) workbook in memory.
    """
    data = get_floor_export_data(floor_name)
    rows = data['rows']
    counts = data['counts']
    date_str = data['date_str']

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{floor_name} Status"

    # Set gridlines visible
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    font_title = Font(name='Segoe UI', size=15, bold=True, color='FFFFFF')
    font_sub = Font(name='Segoe UI', size=9.5, italic=True, color='1E293B')
    font_th = Font(name='Segoe UI', size=10.5, bold=True, color='FFFFFF')
    font_data = Font(name='Segoe UI', size=9.5, color='0F172A')
    font_data_bold = Font(name='Segoe UI', size=9.5, bold=True, color='0F172A')

    fill_banner = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    fill_sub = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    fill_th = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
    fill_alt = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    # 1. Main Title Banner (Row 1)
    ws.merge_cells('A1:H1')
    c1 = ws['A1']
    c1.value = f"ABCD SMART CAMPUS — LIBRARY {floor_name.upper()} DATA REGISTER"
    c1.font = font_title
    c1.fill = fill_banner
    c1.alignment = align_center
    ws.row_dimensions[1].height = 36

    # 2. Subtitle / Timestamp & Summary (Row 2)
    ws.merge_cells('A2:H2')
    c2 = ws['A2']
    c2.value = (
        f"Date: {date_str}   |   Total Seats: {counts['total']}   |   "
        f"Available: {counts['available']}   |   Occupied: {counts['occupied']}   |   "
        f"On Hold: {counts['on_hold']}   |   Shift Occ: {counts['shift_occupied']}   |   Locked: {counts['locked']}"
    )
    c2.font = font_sub
    c2.fill = fill_sub
    c2.alignment = align_center
    ws.row_dimensions[2].height = 24

    # 3. Table Column Headers (Row 4)
    headers = [
        "Seat Number", "Shift / Slot", "Student Name", "Contact Number",
        "Seat Status", "Last Paid Amount", "Fee Expiry Date", "Hold Expiry Date"
    ]
    header_row = 4
    ws.row_dimensions[header_row].height = 26

    for col_idx, text in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.value = text
        cell.font = font_th
        cell.fill = fill_th
        cell.alignment = align_center
        cell.border = border_thin

    # 4. Data Rows
    current_row = 5
    for r in rows:
        ws.row_dimensions[current_row].height = 20
        use_alt = (current_row % 2 == 0)

        data_cells = [
            (f"Seat {r['seat_number']}", align_center, True),
            (r['shift'], align_center, False),
            (r['student_name'], align_left, False),
            (r['mobile'], align_center, False),
            (r['status'], align_center, True),
            (r['last_amount'], align_center, False),
            (r['fee_expiry'], align_center, False),
            (r['hold_expiry'], align_center, False),
        ]

        for col_idx, (val, alignment, is_bold) in enumerate(data_cells, 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = val
            cell.alignment = alignment
            cell.font = font_data_bold if is_bold else font_data
            cell.border = border_thin

            if use_alt:
                cell.fill = fill_alt

            # Status pill highlight
            if col_idx == 5:
                stat = r['status']
                if stat == 'Available':
                    cell.fill = PatternFill('solid', fgColor='DCFCE7')
                    cell.font = Font(name='Segoe UI', size=9.5, bold=True, color='166534')
                elif stat in ('Occupied', 'Shift Occupied'):
                    cell.fill = PatternFill('solid', fgColor='DBEAFE')
                    cell.font = Font(name='Segoe UI', size=9.5, bold=True, color='1E40AF')
                elif stat == 'On Hold':
                    cell.fill = PatternFill('solid', fgColor='FEF3C7')
                    cell.font = Font(name='Segoe UI', size=9.5, bold=True, color='92400E')
                elif stat == 'Temporary':
                    cell.fill = PatternFill('solid', fgColor='FCE7F3')
                    cell.font = Font(name='Segoe UI', size=9.5, bold=True, color='9D174D')
                elif stat == 'Locked':
                    cell.fill = PatternFill('solid', fgColor='F1F5F9')
                    cell.font = Font(name='Segoe UI', size=9.5, bold=True, color='475569')

        current_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip merged banner rows
            if cell.row in (1, 2, 3):
                continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
